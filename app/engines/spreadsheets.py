"""Spreadsheet transformations using openpyxl and standard csv/tsv libraries."""

import csv
import json
from pathlib import Path
from typing import Optional, Dict, Any, List
import openpyxl

from app.engines.base import EngineCapability
from app.core.errors import ConversionFailedError


class SpreadsheetEngine(EngineCapability):
    name = "openpyxl"
    display_name = "Spreadsheet Processor"

    SUPPORTED_INPUTS = ["xlsx", "csv", "tsv"]
    SUPPORTED_OUTPUTS = ["xlsx", "csv", "tsv", "json"]

    @classmethod
    def is_available(cls) -> bool:
        try:
            import openpyxl
            return True
        except ImportError:
            return False

    @classmethod
    def get_version(cls) -> Optional[str]:
        try:
            return f"openpyxl {openpyxl.__version__}"
        except Exception:
            return None

    @classmethod
    def get_binary_path(cls) -> Optional[str]:
        return "bundled-python (openpyxl)"

    @classmethod
    def inspect_spreadsheet(cls, file_path: Path) -> Dict[str, Any]:
        """Inspect sheet names, rows, columns, or CSV line counts."""
        ext = file_path.suffix.lower().lstrip(".")
        info = {
            "format": ext,
            "sheets": [],
            "row_count": 0,
            "column_count": 0
        }
        if ext == "xlsx":
            try:
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                info["sheets"] = wb.sheetnames
                if wb.sheetnames:
                    ws = wb[wb.sheetnames[0]]
                    info["row_count"] = ws.max_row or 0
                    info["column_count"] = ws.max_column or 0
                wb.close()
            except Exception:
                pass
        elif ext in ("csv", "tsv"):
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    lines = [line for line in f]
                    info["row_count"] = len(lines)
                    if lines:
                        delimiter = "\t" if ext == "tsv" else ","
                        reader = csv.reader(lines[:1], delimiter=delimiter)
                        for row in reader:
                            info["column_count"] = len(row)
            except Exception:
                pass
        return info

    @classmethod
    def convert(
        cls,
        source_path: Path,
        target_path: Path,
        target_format: str,
        settings: Dict[str, Any],
        progress_cb=None
    ) -> Path:
        """Convert between XLSX, CSV, TSV, and JSON."""
        src_ext = source_path.suffix.lower().lstrip(".")
        tgt_fmt = target_format.lower()
        target_path.parent.mkdir(parents=True, exist_ok=True)

        encoding = settings.get("encoding", "utf-8")
        delimiter_map = {
            "comma": ",",
            "tab": "\t",
            "semicolon": ";",
            "pipe": "|"
        }
        delimiter = delimiter_map.get(settings.get("delimiter", "comma"), ",")
        if tgt_fmt == "tsv":
            delimiter = "\t"

        quote_map = {
            "minimal": csv.QUOTE_MINIMAL,
            "all": csv.QUOTE_ALL,
            "nonnumeric": csv.QUOTE_NONNUMERIC,
            "none": csv.QUOTE_NONE
        }
        quoting = quote_map.get(settings.get("quoting", "minimal"), csv.QUOTE_MINIMAL)

        if progress_cb:
            progress_cb(10, f"Loading spreadsheet data ({src_ext})...")

        try:
            # 1. XLSX -> CSV / TSV / JSON
            if src_ext == "xlsx":
                wb = openpyxl.load_workbook(source_path, data_only=True)
                sheet_name = settings.get("sheet_name")
                if sheet_name and sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                else:
                    ws = wb.active

                rows_data = []
                for row in ws.iter_rows(values_only=True):
                    # convert None to empty string
                    rows_data.append([str(c) if c is not None else "" for c in row])

                if progress_cb:
                    progress_cb(60, f"Writing output ({tgt_fmt})...")

                if tgt_fmt in ("csv", "tsv"):
                    with open(target_path, "w", encoding=encoding, newline="") as f:
                        writer = csv.writer(f, delimiter=delimiter, quoting=quoting)
                        for r in rows_data:
                            writer.writerow(r)
                elif tgt_fmt == "json":
                    # Convert to array of objects if headers present, or array of arrays
                    if rows_data and len(rows_data) > 1:
                        headers = rows_data[0]
                        records = []
                        for r in rows_data[1:]:
                            record = {headers[i]: r[i] if i < len(r) else "" for i in range(len(headers))}
                            records.append(record)
                        with open(target_path, "w", encoding="utf-8") as f:
                            json.dump(records, f, indent=2)
                    else:
                        with open(target_path, "w", encoding="utf-8") as f:
                            json.dump(rows_data, f, indent=2)

                wb.close()
                if progress_cb:
                    progress_cb(100, "Spreadsheet conversion complete.")
                return target_path

            # 2. CSV / TSV -> XLSX
            elif src_ext in ("csv", "tsv") and tgt_fmt == "xlsx":
                src_delim = "\t" if src_ext == "tsv" else delimiter
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Data"

                with open(source_path, "r", encoding=encoding, errors="replace") as f:
                    reader = csv.reader(f, delimiter=src_delim)
                    for r_idx, row in enumerate(reader, 1):
                        for c_idx, val in enumerate(row, 1):
                            # Try integer / float conversion for proper Excel numeric cells
                            try:
                                if "." in val:
                                    num_val = float(val)
                                else:
                                    num_val = int(val)
                                ws.cell(row=r_idx, column=c_idx, value=num_val)
                            except ValueError:
                                ws.cell(row=r_idx, column=c_idx, value=val)

                if progress_cb:
                    progress_cb(80, "Saving Excel workbook...")

                wb.save(target_path)
                wb.close()
                if progress_cb:
                    progress_cb(100, "Excel workbook generated.")
                return target_path

            # 3. CSV <-> TSV or CSV -> JSON
            elif src_ext in ("csv", "tsv") and tgt_fmt in ("csv", "tsv", "json"):
                src_delim = "\t" if src_ext == "tsv" else ","
                with open(source_path, "r", encoding=encoding, errors="replace") as f:
                    reader = csv.reader(f, delimiter=src_delim)
                    rows_data = list(reader)

                if tgt_fmt in ("csv", "tsv"):
                    with open(target_path, "w", encoding=encoding, newline="") as f:
                        writer = csv.writer(f, delimiter=delimiter, quoting=quoting)
                        for r in rows_data:
                            writer.writerow(r)
                elif tgt_fmt == "json":
                    if rows_data and len(rows_data) > 1:
                        headers = rows_data[0]
                        records = []
                        for r in rows_data[1:]:
                            record = {headers[i]: r[i] if i < len(r) else "" for i in range(len(headers))}
                            records.append(record)
                        with open(target_path, "w", encoding="utf-8") as f:
                            json.dump(records, f, indent=2)
                    else:
                        with open(target_path, "w", encoding="utf-8") as f:
                            json.dump(rows_data, f, indent=2)

                if progress_cb:
                    progress_cb(100, "Delimited data converted.")
                return target_path

            else:
                raise ConversionFailedError(f"Unsupported spreadsheet conversion: {src_ext} to {tgt_fmt}")

        except Exception as e:
            if isinstance(e, ConversionFailedError):
                raise
            raise ConversionFailedError(f"Spreadsheet conversion failed: {str(e)}")
